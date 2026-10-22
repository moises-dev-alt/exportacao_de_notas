from openpyxl import load_workbook

from src.excel_exporter import upsert_excel


def test_add_update_and_ignore_records(tmp_path):
    output = tmp_path / "dados.xlsx"
    first = [{"id": "1", "nome": "Ana", "updated_at": "2026-01-01T10:00:00+00:00"}]
    assert upsert_excel(output, "Dados", first, "id") == {"added": 1, "updated": 0, "ignored": 0}
    assert upsert_excel(output, "Dados", first, "id") == {"added": 0, "updated": 0, "ignored": 1}
    changed = [{"id": "1", "nome": "Ana Maria", "updated_at": "2026-01-02T10:00:00+00:00"}]
    assert upsert_excel(output, "Dados", changed, "id") == {"added": 0, "updated": 1, "ignored": 0}
    workbook = load_workbook(output)
    sheet = workbook["Dados"]
    assert sheet.max_row == 2
    assert sheet["B2"].value == "Ana Maria"
    assert sheet.freeze_panes == "A2"
