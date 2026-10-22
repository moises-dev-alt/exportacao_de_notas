from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from .time_utils import parse_timestamp


def _parse_value(value: Any) -> Any:
    if not isinstance(value, str):
        return value
    try:
        # Excel não armazena timezone. O valor é gravado no horário local de Brasília.
        return parse_timestamp(value).replace(tzinfo=None)
    except ValueError:
        return value


def _columns(existing: list[str], records: list[dict[str, Any]]) -> list[str]:
    result = list(existing)
    for record in records:
        for key in record:
            if key not in result:
                result.append(key)
    return result


def upsert_excel(path: Path, sheet_name: str, records: list[dict[str, Any]], id_field: str) -> dict[str, int]:
    if path.exists():
        workbook = load_workbook(path)
        worksheet = workbook[sheet_name] if sheet_name in workbook.sheetnames else workbook.create_sheet(sheet_name)
    else:
        workbook, worksheet = Workbook(), None
        worksheet = workbook.active
        worksheet.title = sheet_name

    existing_headers = [cell.value for cell in worksheet[1]] if worksheet.max_row else []
    headers = _columns([header for header in existing_headers if header], records)
    if not headers:
        return {"added": 0, "updated": 0, "ignored": 0}
    if id_field not in headers:
        raise ValueError(f"Campo identificador '{id_field}' não existe nos registros.")

    if existing_headers != headers:
        old_rows = list(worksheet.iter_rows(min_row=2, values_only=True)) if existing_headers else []
        old_maps = [dict(zip(existing_headers, row)) for row in old_rows]
        worksheet.delete_rows(1, worksheet.max_row)
        worksheet.append(headers)
        for old_map in old_maps:
            worksheet.append([old_map.get(header) for header in headers])
    elif worksheet.max_row == 0:
        worksheet.append(headers)

    header_index = {header: index + 1 for index, header in enumerate(headers)}
    rows_by_id = {
        worksheet.cell(row=row, column=header_index[id_field]).value: row
        for row in range(2, worksheet.max_row + 1)
        if worksheet.cell(row=row, column=header_index[id_field]).value is not None
    }
    counts = {"added": 0, "updated": 0, "ignored": 0}
    for record in records:
        record_id = record.get(id_field)
        if record_id is None:
            counts["ignored"] += 1
            continue
        values = [_parse_value(record.get(header)) for header in headers]
        row = rows_by_id.get(record_id)
        if row is None:
            worksheet.append(values)
            rows_by_id[record_id] = worksheet.max_row
            counts["added"] += 1
        else:
            old_values = [worksheet.cell(row=row, column=index + 1).value for index in range(len(headers))]
            if old_values == values:
                counts["ignored"] += 1
            else:
                for column, value in enumerate(values, start=1):
                    worksheet.cell(row=row, column=column, value=value)
                counts["updated"] += 1

    _format_sheet(worksheet, headers)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return counts


def _format_sheet(worksheet, headers: list[str]) -> None:
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(1, worksheet.max_row)}"
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for column in range(1, len(headers) + 1):
        letter = get_column_letter(column)
        width = max(len(str(worksheet.cell(row=row, column=column).value or "")) for row in range(1, worksheet.max_row + 1))
        worksheet.column_dimensions[letter].width = min(max(width + 2, 12), 40)
        for row in range(2, worksheet.max_row + 1):
            cell = worksheet.cell(row=row, column=column)
            if isinstance(cell.value, datetime):
                cell.number_format = "dd/mm/yyyy hh:mm"
    for table in list(worksheet.tables.values()):
        del worksheet.tables[table.name]
    if worksheet.max_row >= 2:
        table = Table(displayName="DadosTabela", ref=f"A1:{get_column_letter(len(headers))}{worksheet.max_row}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        worksheet.add_table(table)
